"""EPUB 3, Voxprint Book, and optional FictionBook, DOCX, Markdown, and speaker marks."""

from __future__ import annotations

import base64
import json
import re
import shutil
import uuid
import zipfile
from collections.abc import Collection
from dataclasses import dataclass
from datetime import UTC, datetime
from functools import cache
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from docx import Document
from ebooklib import epub

from plotweaver.appinfo import APP_NAME, APP_VERSION
from plotweaver.domain import ExportFormat, Gender
from plotweaver.errors import assert_never
from plotweaver.paths import repo_root
from plotweaver.vxbook import (
    BlockKind,
    VxBlock,
    VxChapter,
    VxPart,
    VxSpeaker,
    even_utc_now,
    legacy_pair,
    make_character,
    name_key,
    narrator_speaker,
    with_given_aliases,
    write_vxbook,
)

_FB2_NS = "http://www.gribuser.ru/xml/fictionbook/2.0"
_XLINK_NS = "http://www.w3.org/1999/xlink"
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_SPINE_TOC = re.compile(r'\s+toc="ncx"')
_GENRE_ENUM = re.compile(r'<xs:enumeration\s+value="([^"]+)"')
_COVER_ID = "cover.png"
_FALLBACK_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)
# Free-text genre labels the story form uses, mapped onto FictionBook genre codes.
_GENRE_ALIASES = {
    "science fiction": "sf",
    "sci-fi": "sf",
    "scifi": "sf",
    "фантастика": "sf",
    "научная фантастика": "sf",
    "space": "sf_space",
    "space opera": "sf_space",
    "космос": "sf_space",
    "космическая фантастика": "sf_space",
    "horror": "sf_horror",
    "ужасы": "sf_horror",
    "detective": "detective",
    "mystery": "detective",
    "детектив": "detective",
    "thriller": "thriller",
    "триллер": "thriller",
    "romance": "love_contemporary",
    "humor": "humor_prose",
    "humour": "humor_prose",
    "юмор": "humor_prose",
    "history": "prose_history",
    "historical": "prose_history",
    "история": "prose_history",
    "short story": "short_story",
    "adventure": "adventure",
    "приключения": "adventure",
}


@dataclass
class RenderedParagraph:
    text: str
    speaker: str
    gender: Gender
    parts: tuple[tuple[str | None, str], ...] = ()


@dataclass
class RenderedChapter:
    title: str
    paragraphs: list[RenderedParagraph]
    id: str = ""


def speaker_mark(speaker: str, gender: Gender) -> str:
    """One Audiobook Builder mark: NARRATOR, MALE: Name, or FEMALE: Name."""
    if not speaker or speaker.upper() == "NARRATOR":
        return "NARRATOR"
    match gender:
        case Gender.MALE:
            return f"MALE: {speaker}"
        case Gender.FEMALE:
            return f"FEMALE: {speaker}"
        case Gender.UNSPECIFIED:
            return "NARRATOR"
        case _ as unreachable:
            assert_never(unreachable)


def render_markdown(chapters: list[RenderedChapter]) -> tuple[str, str]:
    """Return (book.md, speaker-marks.txt). Marks line up with paragraph blocks, not headings."""
    blocks: list[str] = []
    marks: list[str] = []
    for chapter in chapters:
        title = chapter.title.strip() or "Chapter"
        blocks.append(f"# {title}")
        for paragraph in chapter.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            blocks.append(text)
            marks.append(speaker_mark(paragraph.speaker, paragraph.gender))
    markdown = "\n\n".join(blocks) + "\n"
    marks_text = ("\n".join(marks) + "\n") if marks else ""
    return markdown, marks_text


def write_docx(path: Path, chapters: list[RenderedChapter]) -> None:
    document = Document()
    for chapter in chapters:
        document.add_heading(chapter.title.strip() or "Chapter", level=1)
        for paragraph in chapter.paragraphs:
            text = paragraph.text.strip()
            if text:
                document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def fb2_schema_dir() -> Path:
    """Vendored FictionBook 2.1 schema (BSD-2-Clause, Dmitry Gribov)."""
    return Path(__file__).resolve().parent / "data" / "fb2"


@cache
def genre_codes() -> frozenset[str]:
    text = (fb2_schema_dir() / "FictionBookGenres.xsd").read_text(encoding="utf-8")
    return frozenset(_GENRE_ENUM.findall(text))


def fb2_genre(raw: str, language: str) -> str:
    """Map a free-text genre onto a code the FictionBook schema accepts."""
    codes = genre_codes()
    key = " ".join(raw.replace("_", " ").split()).strip().lower()
    underscored = key.replace(" ", "_")
    if underscored in codes:
        return underscored
    if key in {"fantasy", "фэнтези"}:
        code = "russian_fantasy" if language == "ru" else "foreign_fantasy"
        if code in codes:
            return code
    mapped = _GENRE_ALIASES.get(key, "")
    if mapped in codes:
        return mapped
    for token, code in (
        ("sf_space", "sf_space"),
        ("space", "sf_space"),
        ("космос", "sf_space"),
        ("марс", "sf_space"),
        ("mars", "sf_space"),
        ("horror", "sf_horror"),
        ("ужас", "sf_horror"),
        ("detective", "detective"),
        ("детектив", "detective"),
        ("thriller", "thriller"),
        ("fantasy", "russian_fantasy" if language == "ru" else "foreign_fantasy"),
        ("фэнтези", "russian_fantasy"),
        ("science", "sf"),
        ("фантаст", "sf"),
        ("humor", "humor_prose"),
        ("юмор", "humor_prose"),
        ("romance", "love_contemporary"),
        ("histori", "prose_history"),
        ("истор", "prose_history"),
    ):
        if token in key and code in codes:
            return code
    if "narrative" in codes:
        return "narrative"
    return sorted(codes)[0]


def cover_png() -> bytes:
    """Suite mark used as the book cover. A 1×1 PNG is the last resort."""
    candidates: list[Path] = []
    try:
        candidates.append(repo_root() / "branding" / "plotweaver-512.png")
    except FileNotFoundError:
        pass
    candidates.append(Path(__file__).resolve().parents[2] / "branding" / "plotweaver-512.png")
    for path in candidates:
        if path.is_file():
            return path.read_bytes()
    return _FALLBACK_PNG


def write_fb2(
    path: Path,
    *,
    title: str,
    language: str,
    author: str,
    genre: str,
    synopsis: str,
    chapters: list[RenderedChapter],
    cover: bytes | None = None,
) -> None:
    """FictionBook 2.1: title-info, cover, annotation, and a titled section per chapter."""
    ET.register_namespace("", _FB2_NS)
    ET.register_namespace("l", _XLINK_NS)
    root = ET.Element(_tag("FictionBook"))
    description = ET.SubElement(root, _tag("description"))
    title_info = ET.SubElement(description, _tag("title-info"))
    ET.SubElement(title_info, _tag("genre")).text = fb2_genre(genre, language)
    _append_author(title_info, author)
    ET.SubElement(title_info, _tag("book-title")).text = _xml(title) or "Untitled"
    paragraphs = _annotation_paragraphs(synopsis)
    if paragraphs:
        annotation = ET.SubElement(title_info, _tag("annotation"))
        for paragraph in paragraphs:
            ET.SubElement(annotation, _tag("p")).text = paragraph
    coverpage = ET.SubElement(title_info, _tag("coverpage"))
    image = ET.SubElement(coverpage, _tag("image"))
    image.set(_xlink("type"), "simple")
    image.set(_xlink("href"), f"#{_COVER_ID}")
    ET.SubElement(title_info, _tag("lang")).text = language or "ru"

    document = ET.SubElement(description, _tag("document-info"))
    maker = ET.SubElement(document, _tag("author"))
    ET.SubElement(maker, _tag("nickname")).text = APP_NAME
    ET.SubElement(document, _tag("program-used")).text = f"{APP_NAME} {APP_VERSION}"
    today = datetime.now(UTC).date().isoformat()
    date = ET.SubElement(document, _tag("date"))
    date.set("value", today)
    date.text = today
    ET.SubElement(document, _tag("id")).text = str(uuid.uuid4())
    ET.SubElement(document, _tag("version")).text = "1.0"

    body = ET.SubElement(root, _tag("body"))
    if not chapters:
        chapters = [RenderedChapter("Chapter", [])]
    for chapter in chapters:
        section = ET.SubElement(body, _tag("section"))
        heading = ET.SubElement(section, _tag("title"))
        ET.SubElement(heading, _tag("p")).text = _xml(chapter.title) or "Chapter"
        texts = [_xml(paragraph.text) for paragraph in chapter.paragraphs if paragraph.text.strip()]
        if texts:
            for text in texts:
                ET.SubElement(section, _tag("p")).text = text
        else:
            ET.SubElement(section, _tag("empty-line"))

    binary = ET.SubElement(root, _tag("binary"))
    binary.set("id", _COVER_ID)
    binary.set("content-type", "image/png")
    encoded = base64.b64encode(cover if cover is not None else cover_png()).decode("ascii")
    binary.text = "\n".join(encoded[index : index + 76] for index in range(0, len(encoded), 76))

    ET.indent(root, space="  ")
    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    if not payload.endswith(b"\n"):
        payload += b"\n"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)


def write_epub(
    path: Path,
    title: str,
    language: str,
    chapters: list[RenderedChapter],
    *,
    author: str = "",
    cover: bytes | None = None,
) -> None:
    """EPUB 3 package: navigation document, cover, and one XHTML file per chapter."""
    book = epub.EpubBook()
    book.set_identifier(f"urn:uuid:{uuid.uuid4()}")
    book.set_title(_xml(title) or "Untitled")
    book.set_language(language or "en")
    book.add_author(_author_label(author))
    image = cover if cover is not None else cover_png()
    book.set_cover(_COVER_ID, image, create_page=True)
    items: list[epub.EpubHtml] = []
    if not chapters:
        chapters = [RenderedChapter("Chapter", [])]
    for index, chapter in enumerate(chapters, start=1):
        heading = _xml(chapter.title) or "Chapter"
        body = [f"<h1>{escape(heading)}</h1>"]
        for paragraph in chapter.paragraphs:
            text = paragraph.text.strip()
            if text:
                body.append(f"<p>{escape(_xml(text))}</p>")
        item = epub.EpubHtml(title=heading, file_name=f"ch{index:02d}.xhtml", lang=language or "en")
        item.content = "".join(body)
        book.add_item(item)
        items.append(item)
    book.toc = tuple(items)
    book.add_item(epub.EpubNav())
    guide: list[dict[str, str]] = [{"href": "cover.xhtml", "type": "cover", "title": "Cover"}]
    guide.append({"href": items[0].file_name, "type": "text", "title": items[0].title or "Start"})
    book.guide = guide
    book.spine = ["cover", "nav", *items]
    path.parent.mkdir(parents=True, exist_ok=True)
    options = {
        "epub2_guide": False,
        "epub3_pages": False,
        "epub3_landmark": True,
        "raise_exceptions": True,
    }
    written = epub.write_epub(str(path), book, options)
    if written is False:
        raise OSError(f"Could not write {path}")
    _patch_epub(path)


def write_bundle(
    directory: Path,
    title: str,
    language: str,
    chapters: list[RenderedChapter],
    overlap: dict[str, object],
    *,
    author: str = "",
    genre: str = "",
    synopsis: str = "",
    formats: Collection[str] | None = None,
    created: str = "",
    speakers: list[VxSpeaker] | None = None,
) -> tuple[str, str]:
    """Write book.vxbook and EPUB 3, plus any extra formats the user asked for."""
    directory.mkdir(parents=True, exist_ok=True)
    stamp = created or even_utc_now()
    vx_chapters, vx_speakers = bind_rendered(chapters, speakers or [])
    write_vxbook(
        directory / "book.vxbook",
        title=title,
        author=author,
        language=language,
        chapters=vx_chapters,
        speakers=vx_speakers,
        created=stamp,
        synopsis=synopsis,
    )
    markdown, _plain_marks = render_markdown(chapters)
    selected = _selected_formats(formats)
    marks = ""
    for kind in selected:
        match kind:
            case ExportFormat.EPUB:
                write_epub(directory / "book.epub", title, language, chapters, author=author)
            case ExportFormat.FB2:
                write_fb2(
                    directory / "book.fb2",
                    title=title,
                    language=language,
                    author=author,
                    genre=genre,
                    synopsis=synopsis,
                    chapters=chapters,
                )
            case ExportFormat.DOCX:
                write_docx(directory / "book.docx", chapters)
            case ExportFormat.MARKDOWN:
                (directory / "book.md").write_text(markdown, encoding="utf-8", newline="\n")
            case ExportFormat.SPEAKER_MARKS:
                plain, marks = legacy_pair(directory / "book.vxbook")
                audio = directory / "audiobook"
                audio.mkdir(exist_ok=True)
                (audio / "book.md").write_text(plain, encoding="utf-8", newline="\n")
                (audio / "speaker-marks.txt").write_text(marks, encoding="utf-8", newline="\n")
            case _ as unreachable:
                assert_never(unreachable)
    payload = json.dumps(overlap, ensure_ascii=False, indent=2) + "\n"
    (directory / "overlap-report.json").write_text(payload, encoding="utf-8", newline="\n")
    return markdown, marks


def publish_exports(source: Path, dest: Path) -> None:
    """Copy the export bundle to the folder the user asked for."""
    if source.resolve() == dest.resolve():
        return
    dest.mkdir(parents=True, exist_ok=True)
    names = ("book.md", "book.docx", "book.epub", "book.fb2", "book.vxbook", "overlap-report.json")
    for name in names:
        item = source / name
        if item.is_file():
            shutil.copy2(item, dest / name)
    audio = source / "audiobook"
    if audio.is_dir():
        target = dest / "audiobook"
        target.mkdir(exist_ok=True)
        for item in audio.iterdir():
            if item.is_file():
                shutil.copy2(item, target / item.name)


def bind_rendered(
    chapters: list[RenderedChapter], cards: list[VxSpeaker]
) -> tuple[list[VxChapter], list[VxSpeaker]]:
    """Map canon speaker names onto vxbook ids. Inferred speech is src=auto."""
    people = [narrator_speaker(), *cards]
    by_key: dict[str, VxSpeaker] = {}
    for speaker in people:
        if speaker.name:
            by_key.setdefault(name_key(speaker.name), speaker)
        for alias in speaker.aliases:
            by_key.setdefault(name_key(alias), speaker)
    used = {speaker.id for speaker in people}
    extras: list[VxSpeaker] = []

    def resolve(name: str, gender: Gender) -> str | None:
        key = name_key(name)
        if not key or key == "narrator" or name.strip().upper() == "NARRATOR":
            return None
        found = by_key.get(key)
        if found is not None:
            return found.id
        created = make_character(name, gender.value, used)
        extras.append(created)
        by_key[name_key(created.name)] = created
        return created.id

    vx_chapters: list[VxChapter] = []
    for index, chapter in enumerate(chapters, start=1):
        blocks: list[VxBlock] = []
        for paragraph in chapter.paragraphs:
            parts = tuple(
                VxPart(
                    text,
                    resolve(speaker, Gender.UNSPECIFIED) if speaker else None,
                    "auto" if speaker else None,
                )
                for speaker, text in paragraph.parts
            )
            if not paragraph.text.strip() and not parts:
                continue
            speaker_id = resolve(paragraph.speaker, paragraph.gender)
            blocks.append(
                VxBlock(
                    BlockKind.PARAGRAPH,
                    text=paragraph.text,
                    speaker_id=speaker_id,
                    src="auto" if speaker_id else None,
                    parts=parts,
                )
            )
        vx_chapters.append(VxChapter(chapter.id or f"ch{index:02d}", chapter.title, blocks))
    return vx_chapters, with_given_aliases([*people, *extras])


def _selected_formats(formats: Collection[str] | None) -> list[ExportFormat]:
    names = ["epub"]
    if formats:
        names.extend(formats)
    chosen: list[ExportFormat] = []
    seen: set[str] = set()
    allowed = {item.value for item in ExportFormat}
    for name in names:
        if name in allowed and name not in seen:
            seen.add(name)
            chosen.append(ExportFormat(name))
    return chosen


def _tag(name: str) -> str:
    return f"{{{_FB2_NS}}}{name}"


def _xlink(name: str) -> str:
    return f"{{{_XLINK_NS}}}{name}"


def _xml(text: str) -> str:
    return _ILLEGAL_XML.sub("", text).strip()


def _author_label(author: str) -> str:
    return _xml(author) or "Plotweaver"


def _append_author(parent: ET.Element, author: str) -> None:
    node = ET.SubElement(parent, _tag("author"))
    words = _author_label(author).split()
    if len(words) == 2:
        ET.SubElement(node, _tag("first-name")).text = words[0]
        ET.SubElement(node, _tag("last-name")).text = words[1]
        return
    if len(words) == 3:
        ET.SubElement(node, _tag("first-name")).text = words[0]
        ET.SubElement(node, _tag("middle-name")).text = words[1]
        ET.SubElement(node, _tag("last-name")).text = words[2]
        return
    ET.SubElement(node, _tag("nickname")).text = " ".join(words) or "Plotweaver"


def _annotation_paragraphs(synopsis: str) -> list[str]:
    text = synopsis.replace("\r\n", "\n").replace("\r", "\n").strip()
    paragraphs: list[str] = []
    for block in text.split("\n\n"):
        lines: list[str] = []
        for line in block.split("\n"):
            stripped = line.strip()
            if stripped.startswith("#"):
                stripped = stripped.lstrip("#").strip()
            if stripped:
                lines.append(stripped)
        if lines:
            paragraphs.append(_xml(" ".join(lines)))
    return [paragraph for paragraph in paragraphs if paragraph]


def _patch_epub(path: Path) -> None:
    """Remove EPUB 2 leftovers ebooklib still writes into an EPUB 3 package."""
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    tmp = path.with_suffix(".epub.tmp")
    with zipfile.ZipFile(tmp, "w") as target:
        for info, data in entries:
            if info.filename.endswith((".xhtml", ".html", ".opf")):
                text = data.decode("utf-8")
                text = text.replace(' role="doc-toc"', "")
                text = text.replace(
                    ' epub:prefix="z3998: http://www.daisy.org/z3998/2012/vocab/structure/#"',
                    "",
                )
                text = _SPINE_TOC.sub("", text)
                data = text.encode("utf-8")
            cloned = zipfile.ZipInfo(filename=info.filename, date_time=info.date_time)
            cloned.compress_type = info.compress_type
            cloned.external_attr = info.external_attr
            cloned.flag_bits = info.flag_bits
            target.writestr(cloned, data)
    tmp.replace(path)
